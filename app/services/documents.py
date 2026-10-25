from pathlib import Path
from uuid import uuid4
from docx import Document as DocxDocument
from pypdf import PdfReader
from fastapi import HTTPException, UploadFile
from app.core.config import get_settings

ALLOWED = {"application/pdf": ".pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx", "text/plain": ".txt"}
async def save_and_extract(file: UploadFile) -> tuple[Path, str, dict[str, str]]:
    if file.content_type not in ALLOWED: raise HTTPException(415, "Only PDF, DOCX and TXT documents are accepted")
    raw = await file.read()
    if len(raw) > 15 * 1024 * 1024: raise HTTPException(413, "Maximum file size is 15 MB")
    settings = get_settings(); settings.upload_dir.mkdir(parents=True, exist_ok=True)
    path = settings.upload_dir / f"{uuid4()}{ALLOWED[file.content_type]}"; path.write_bytes(raw)
    try:
        if file.content_type == "application/pdf":
            reader=PdfReader(str(path)); content="\n".join(page.extract_text() or "" for page in reader.pages); meta={"pages":str(len(reader.pages))}
        elif file.content_type.endswith("document"):
            doc=DocxDocument(str(path)); content="\n".join(p.text for p in doc.paragraphs); meta={"paragraphs":str(len(doc.paragraphs))}
        else: content=raw.decode("utf-8", errors="replace"); meta={"encoding":"utf-8"}
    except Exception as exc:
        path.unlink(missing_ok=True); raise HTTPException(422, "Document cannot be parsed") from exc
    return path, content, meta
